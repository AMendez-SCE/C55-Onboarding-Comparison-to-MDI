import pandas as pd
import numpy as np
import re
from pathlib import Path
from difflib import SequenceMatcher
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule

# =============================================================================
# FILE PATHS
# =============================================================================

base_path = r"C:\Users\mendezal\Documents\PDM\TPP Project Migration Analysis"

input_file = Path(base_path) / "MDI TPP List Analysis 10.12.26.xlsx"
output_file = Path(base_path) / "MDI_C55_Reconciliation_for_TPP.xlsx"

# =============================================================================
# LOAD DATA
# =============================================================================

mdi_df = pd.read_excel(
    input_file,
    sheet_name="MDI Migrated Projects",
    engine="openpyxl"
)

c55_df = pd.read_excel(
    input_file,
    sheet_name="C55 TPPs",
    engine="openpyxl"
)

mdi_key = "Shared Cost Number"
c55_key = "Legacy SCN"

# =============================================================================
# NORMALIZATION FUNCTIONS
# =============================================================================

def normalize_column_name(col):
    if pd.isna(col):
        return ""

    col = str(col).lower().strip()

    replacements = {
        "_": "",
        "-": "",
        " ": "",
        "/": "",
        ".": ""
    }

    for old, new in replacements.items():
        col = col.replace(old, new)

    return col

def normalize_value(value):

    if pd.isna(value):
        return ""

    # Date handling
    try:
        dt = pd.to_datetime(value, errors="raise")
        return dt.strftime("%Y-%m-%d")
    except:
        pass

    # Currency / numeric handling
    try:
        if isinstance(value, str):
            temp = value.replace("$", "").replace(",", "").strip()
            numeric = float(temp)
            return str(round(numeric))
        elif isinstance(value, (int, float, np.number)):
            return str(round(float(value)))
    except:
        pass

    return str(value).strip()

# =============================================================================
# CLEAN KEYS
# =============================================================================

mdi_df[mdi_key] = mdi_df[mdi_key].astype(str).str.strip()
c55_df[c55_key] = c55_df[c55_key].astype(str).str.strip()

# =============================================================================
# MATCH COLUMNS
# =============================================================================

mdi_cols = list(mdi_df.columns)
c55_cols = list(c55_df.columns)

column_mapping = {}

for mdi_col in mdi_cols:

    if mdi_col == mdi_key:
        continue

    best_match = None
    best_score = 0

    mdi_clean = normalize_column_name(mdi_col)

    for c55_col in c55_cols:

        if c55_col == c55_key:
            continue

        c55_clean = normalize_column_name(c55_col)

        score = SequenceMatcher(
            None,
            mdi_clean,
            c55_clean
        ).ratio()

        if score > best_score:
            best_score = score
            best_match = c55_col

    if best_score >= 0.65:
        column_mapping[mdi_col] = best_match

# =============================================================================
# BUILD LOOKUPS
# =============================================================================
mdi_df = mdi_df.dropna(subset=[mdi_key])
c55_df = c55_df.dropna(subset=[c55_key])

mdi_df = mdi_df.drop_duplicates(subset=[mdi_key], keep="last")
c55_df = c55_df.drop_duplicates(subset=[c55_key], keep="last")

mdi_lookup = mdi_df.set_index(mdi_key).to_dict('index')
c55_lookup = c55_df.set_index(c55_key).to_dict('index')

all_scns = sorted(
    set(mdi_lookup.keys()).union(set(c55_lookup.keys()))
)

# =============================================================================
# RECONCILIATION
# =============================================================================

results = []

for scn in all_scns:

    mdi_row = mdi_lookup.get(scn)
    c55_row = c55_lookup.get(scn)

    if mdi_row is None:

        results.append({
            "Shared Cost Number": scn,
            "Column": "",
            "MDI Value": "",
            "C55 Value": "",
            "Status": "Missing in MDI"
        })
        continue

    if c55_row is None:

        results.append({
            "Shared Cost Number": scn,
            "Column": "",
            "MDI Value": "",
            "C55 Value": "",
            "Status": "Missing in C55"
        })
        continue

    for mdi_col, c55_col in column_mapping.items():

        mdi_value = normalize_value(
            mdi_row.get(mdi_col)
        )

        c55_value = normalize_value(
            c55_row.get(c55_col)
        )

        status = (
            "MATCH"
            if mdi_value == c55_value
            else "DIFFERENT"
        )

        results.append({
            "Shared Cost Number": scn,
            "MDI Column": mdi_col,
            "C55 Column": c55_col,
            "MDI Value": mdi_value,
            "C55 Value": c55_value,
            "Status": status
        })

# =============================================================================
# RESULTS DATAFRAME
# =============================================================================

results_df = pd.DataFrame(results)

summary_df = (
    results_df["Status"]
    .value_counts()
    .reset_index()
)

summary_df.columns = [
    "Status",
    "Count"
]

mapping_df = pd.DataFrame(
    list(column_mapping.items()),
    columns=["MDI Column", "C55 Column"]
)

# =============================================================================
# WRITE OUTPUT
# =============================================================================

with pd.ExcelWriter(
    output_file,
    engine="openpyxl"
) as writer:

    results_df.to_excel(
        writer,
        sheet_name="Reconciliation",
        index=False
    )

    summary_df.to_excel(
        writer,
        sheet_name="Summary",
        index=False
    )

    mapping_df.to_excel(
        writer,
        sheet_name="Column Mapping",
        index=False
    )

# =============================================================================
# FORMATTING
# =============================================================================

wb = load_workbook(output_file)

recon_ws = wb["Reconciliation"]

header_fill = PatternFill(
    start_color="D9D9D9",
    end_color="D9D9D9",
    fill_type="solid"
)

red_fill = PatternFill(
    start_color="FFC7CE",
    end_color="FFC7CE",
    fill_type="solid"
)

yellow_fill = PatternFill(
    start_color="FFF2CC",
    end_color="FFF2CC",
    fill_type="solid"
)

# Header formatting
for cell in recon_ws[1]:
    cell.font = Font(bold=True)
    cell.fill = header_fill

# Find Status column
status_col = None

for cell in recon_ws[1]:
    if str(cell.value).strip().upper() == "STATUS":
        status_col = cell.column
        break

# Highlight rows
if status_col:

    status_letter = get_column_letter(status_col)

    red_fill = PatternFill(
        fill_type="solid",
        start_color="FFC7CE",
        end_color="FFC7CE"
    )

    yellow_fill = PatternFill(
        fill_type="solid",
        start_color="FFF2CC",
        end_color="FFF2CC"
    )

    # Red rows for DIFFERENT
    recon_ws.conditional_formatting.add(
        f"A2:{get_column_letter(recon_ws.max_column)}{recon_ws.max_row}",
        FormulaRule(
            formula=[
                f'${status_letter}2="DIFFERENT"'
            ],
            fill=red_fill
        )
    )

    # Yellow rows for Missing records
    recon_ws.conditional_formatting.add(
        f"A2:{get_column_letter(recon_ws.max_column)}{recon_ws.max_row}",
        FormulaRule(
            formula=[
                f'OR(${status_letter}2="Missing in MDI",${status_letter}2="Missing in C55")'
            ],
            fill=yellow_fill
        )
    )
    

for column in recon_ws.columns:

    max_length = 0

    col_letter = get_column_letter(
        column[0].column
    )

    for cell in column:
        try:
            max_length = max(
                max_length,
                len(str(cell.value))
            )
        except:
            pass

    recon_ws.column_dimensions[
        col_letter
    ].width = min(max_length + 3, 60)

for ws in wb.worksheets:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

wb.save(output_file)

print()
print("=" * 60)
print("RECONCILIATION COMPLETE")
print("=" * 60)
print(f"Input File : {input_file}")
print(f"Output File: {output_file}")
print("=" * 60)